from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import event, text
from werkzeug.datastructures import FileStorage

from app.extensions import db
from app.models.db import (
    CategoriaProduto, Funcionario, FuncionarioEmpresa, ItemVenda, LancamentoFinanceiro,
    MovimentoEstoque, OperacaoIdempotente, Produto, ProdutoEmpresa, Role, RolePermission,
    Tenant, TipoMovimentoEstoque, MotivoMovimentoEstoque, Venda,
)
from app.services.acesso_empresa_service import AcessoEmpresaService
from app.services.categoria_service import CategoriaService
from app.services.estoque_service import EstoqueService
from app.services.funcionario_service import FuncionarioService
from app.services.import_export_service import ImportExportService
from app.services.pdv_service import PdvService
from app.services.produto_service import ProdutoService
from app.services.role_service import RoleService
from app.services.tenant_bootstrap_service import TenantBootstrapService
from tests.test_sprint02_security import PASSWORD, security_app, login, csrf_headers


@pytest.fixture
def transaction_app(security_app):
    with security_app.app_context():
        TenantBootstrapService.garantir_cadastros_operacionais(1)
        db.session.commit()
    return security_app


def scope():
    return AcessoEmpresaService.obter_escopo(1, 1)


def product(name="Produto", **fields):
    return ProdutoService.criar({"nome": name, "empresa_id": 1, "valor_varejo": "10",
                                 "valor_atacado": "8", **fields}, 1, scope(), 1)


def employee(**fields):
    role = Role.query.filter_by(tenant_id=1, codigo="operador").one()
    return {"nome": "Equipe", "cpf": "99988877766", "usuario": "equipe", "senha": PASSWORD,
            "empresa_id": 1, "empresa_ids": [1, 2], "role_id": role.id, **fields}


def upload(entity, rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "dados"
    columns = headers or [column["key"] for column in ImportExportService.ENTITY_DEFINITIONS[entity]["columns"]]
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column) for column in columns])
    output = BytesIO()
    workbook.save(output)
    return FileStorage(stream=BytesIO(output.getvalue()), filename="lote.xlsx")


def import_rows(entity, rows, **options):
    return ImportExportService.importar_entidade(entity, upload(entity, rows), 1, scope(), 1, **options)


def parallel(app, action, count=2):
    barrier = Barrier(count)

    def run(index):
        with app.app_context():
            db.session.execute(text("SET statement_timeout = '15s'"))
            db.session.commit()
            barrier.wait(timeout=10)
            try:
                return action(index)
            except ValueError as error:
                return str(error)
            finally:
                db.session.remove()

    with ThreadPoolExecutor(max_workers=count) as executor:
        return list(executor.map(run, range(count)))


def stock_setup(quantity=10):
    record = product()
    record.estoque_atual = quantity
    db.session.commit()
    return record.id, record.produto_id


def sale_payload(product_id, quantity=1, key="sale-1"):
    from app.models.db import FormaPagamento
    payment = FormaPagamento.query.filter_by(tenant_id=1, nome="Dinheiro").one()
    return {"empresa_id": 1, "itens": [{"produto_id": product_id, "quantidade": quantity}],
            "pagamentos": [{"forma_pagamento_id": payment.id, "valor": str(10 * quantity)}],
            "idempotency_key": key}


@pytest.mark.parametrize("model,operation", [
    (ProdutoEmpresa, "product"), (FuncionarioEmpresa, "employee"), (RolePermission, "role"),
])
def test_parent_and_links_rollback_after_child_insert_failure(transaction_app, model, operation):
    with transaction_app.app_context():
        before = {table: db.session.query(table).count() for table in (Produto, ProdutoEmpresa, Funcionario, FuncionarioEmpresa, Role, RolePermission)}

        def fail(*args):
            raise RuntimeError("injected child failure")

        event.listen(model, "after_insert", fail)
        try:
            with pytest.raises(RuntimeError, match="injected"):
                if operation == "product":
                    product()
                elif operation == "employee":
                    FuncionarioService.criar(employee(), 1)
                else:
                    permission_id = RolePermission.query.filter_by(tenant_id=1).first().permission_id
                    RoleService.criar({"nome": "Novo", "codigo": "novo", "permission_ids": [permission_id]}, 1)
        finally:
            event.remove(model, "after_insert", fail)
        assert {table: db.session.query(table).count() for table in before} == before
        assert not db.session.new and not db.session.dirty


@pytest.mark.parametrize("entity", ["product", "employee"])
def test_final_parent_delete_failure_restores_link(transaction_app, entity):
    with transaction_app.app_context():
        if entity == "product":
            record = product()
            model, link_model, remove = Produto, ProdutoEmpresa, lambda record_id: ProdutoService.deletar(record_id, 1, scope())
        else:
            record = FuncionarioService.criar(employee(empresa_ids=[1]), 1)
            model, link_model, remove = Funcionario, FuncionarioEmpresa, lambda record_id: FuncionarioService.deletar(record_id, 1)
        record_id = record.id
        total = db.session.query(model).count()

        def fail(*args):
            raise RuntimeError("injected parent failure")

        event.listen(model, "after_delete", fail)
        try:
            with pytest.raises(RuntimeError, match="injected"):
                remove(record_id)
        finally:
            event.remove(model, "after_delete", fail)
        assert db.session.get(link_model, record_id) is not None
        assert db.session.query(model).count() == total


def test_employee_company_set_and_role_changes_are_atomic(transaction_app):
    with transaction_app.app_context():
        record = FuncionarioService.criar(employee(), 1)
        employee_id, record_id = record.funcionario_id, record.id
        assert FuncionarioEmpresa.query.filter_by(funcionario_id=employee_id).count() == 2
        with pytest.raises(ValueError, match="Empresa"):
            FuncionarioService.atualizar(record_id, employee(nome="Nao persistir", empresa_ids=[1, 3]), 1)
        assert db.session.get(Funcionario, employee_id).nome == "Equipe"
        assert FuncionarioEmpresa.query.filter_by(funcionario_id=employee_id).count() == 2
        result = FuncionarioService.atualizar(record_id, employee(empresa_id=2, empresa_ids=[2]), 1)
        assert result.empresa_id == 2
        assert FuncionarioEmpresa.query.filter_by(funcionario_id=employee_id).count() == 1


def test_product_multiple_companies_preserves_balances_and_quota(transaction_app):
    with transaction_app.app_context():
        first = product()
        product_id, first_id = first.produto_id, first.id
        db.session.get(Tenant, 1).limite_produtos = 1
        db.session.commit()
        second = ProdutoService.criar({"produto_id": product_id, "empresa_id": 2, "valor_varejo": "20"}, 1, scope(), 1)
        second_id = second.id
        assert second.estoque_atual == 0 and Produto.query.count() == 1
        with pytest.raises(ValueError, match="vinculado"):
            ProdutoService.criar({"produto_id": product_id, "empresa_id": 2}, 1, scope(), 1)
        ProdutoService.deletar(first_id, 1, scope())
        assert db.session.get(Produto, product_id) is not None
        ProdutoService.deletar(second_id, 1, scope())
        assert Produto.query.count() == ProdutoEmpresa.query.count() == 0


@pytest.mark.parametrize("fields", [
    {"nome": " "}, {"nome": "a" * 151}, {"valor_varejo": "NaN"},
    {"valor_atacado": "11"}, {"quantidade_minima_atacado": "1.5"},
    {"estoque_minimo": "1,5"}, {"codigo_barras": "ABC123"}, {"categoria_id": 999},
    {"empresa_id": 3}, {"possui_ncm": True, "ncm": ""},
])
def test_invalid_product_never_leaves_parent(transaction_app, fields):
    with transaction_app.app_context():
        with pytest.raises((ValueError, PermissionError)):
            product(**fields)
        assert Produto.query.count() == ProdutoEmpresa.query.count() == 0


def test_category_validation_and_in_use_delete(transaction_app):
    with transaction_app.app_context():
        with pytest.raises(ValueError):
            CategoriaService.criar({"nome": " "}, 1)
        category = CategoriaService.criar({"nome": "Categoria"}, 1)
        category_id = category.id
        product(categoria_id=category_id)
        with pytest.raises(ValueError):
            CategoriaService.criar({"nome": "Categoria"}, 1)
        with pytest.raises(ValueError, match="vinculada"):
            CategoriaService.deletar(category_id, 1)
        assert db.session.get(CategoriaProduto, category_id) is not None


def test_concurrent_generated_and_duplicate_barcodes(transaction_app):
    results = parallel(transaction_app, lambda index: product(f"Generated {index}").produto.codigo_barras, 4)
    assert len(set(results)) == 4 and all(len(code) == 13 for code in results)
    results = parallel(transaction_app, lambda index: product(f"Duplicate {index}", codigo_barras="123456789").id)
    assert sum(isinstance(result, int) for result in results) == 1
    with transaction_app.app_context():
        assert Produto.query.count() == ProdutoEmpresa.query.count() == 5


def test_concurrent_stock_exhaustion_has_one_winner(transaction_app):
    with transaction_app.app_context():
        record_id, _ = stock_setup(1)

    def remove(index):
        result = EstoqueService.registrar_movimentacao_manual({"produto_empresa_id": record_id,
            "tipo_movimento": "SAIDA", "motivo": "AJUSTE", "quantidade": 1}, 1, scope(), 1)
        return result.id

    results = parallel(transaction_app, remove, 4)
    assert sum(isinstance(result, int) for result in results) == 1
    assert sum("insuficiente" in str(result) for result in results) == 3
    with transaction_app.app_context():
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 0
        assert MovimentoEstoque.query.count() == 1


def test_sale_and_partial_cancellation_replay_concurrently(transaction_app):
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        payload = sale_payload(product_id, quantity=2)
    results = parallel(transaction_app, lambda index: PdvService.criar_venda(payload, 1, scope(), 1), 3)
    assert len({result["id"] for result in results}) == 1
    sale_id = results[0]["id"]
    item_id = results[0]["itens"][0]["id"]
    with transaction_app.app_context():
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 8
        assert Venda.query.count() == ItemVenda.query.count() == MovimentoEstoque.query.count() == 1
        with pytest.raises(ValueError, match="outros dados"):
            PdvService.criar_venda({**payload, "observacao": "changed"}, 1, scope(), 1)
    cancelled = parallel(transaction_app, lambda index: PdvService.cancelar_item_venda(
        sale_id, item_id, {"quantidade": 1, "idempotency_key": "partial-1"}, 1, scope(), 1), 3)
    assert all(result["itens"][0]["quantidade_cancelada"] == 1 for result in cancelled)
    with transaction_app.app_context():
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 9
        assert MovimentoEstoque.query.count() == 2
        assert OperacaoIdempotente.query.count() == 2


def test_full_cancellation_and_stock_callback_replays(transaction_app):
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        sale = PdvService.criar_venda(sale_payload(product_id), 1, scope(), 1)
        sale_id = sale["id"]
        item = sale["itens"][0]
        items = [{"item_venda_id": item["id"], "produto_id": product_id, "quantidade": 1}]
        EstoqueService.registrar_saida_por_venda(sale_id, 1, items, 1, 1, scope())
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 9
    results = parallel(transaction_app, lambda index: PdvService.cancelar_venda(sale_id, {}, 1, scope(), 1), 3)
    assert all(result["status"] == "CANCELADA" for result in results)
    with transaction_app.app_context():
        EstoqueService.registrar_entrada_por_cancelamento_venda(sale_id, 1, items, 1, 1, scope())
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 10
        assert MovimentoEstoque.query.count() == 2


def test_sale_failure_after_stock_restores_every_table(transaction_app, monkeypatch):
    from app.services.financeiro_service import FinanceiroService
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        payload = sale_payload(product_id)

        def fail(**kwargs):
            db.session.flush()
            raise RuntimeError("injected finance failure")

        monkeypatch.setattr(FinanceiroService, "registrar_entradas_da_venda", fail)
        with pytest.raises(RuntimeError, match="injected"):
            PdvService.criar_venda(payload, 1, scope(), 1)
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 10
        for model in (Venda, ItemVenda, MovimentoEstoque, LancamentoFinanceiro, OperacaoIdempotente):
            assert db.session.query(model).count() == 0


LAYOUT_ROWS = {
    "categorias": {"nome": "Importada", "descricao": "Descricao"},
    "produtos": {"empresa": "Empresa 1-1", "nome": "Importado", "estoque_atual": 3,
                 "valor_varejo": "10", "valor_atacado": "8", "quantidade_minima_atacado": 5},
    "funcionarios": {"empresa": "Empresa 1-1", "role": "operador", "nome": "Equipe",
                     "cpf": "22233344455", "usuario": "importado", "senha": PASSWORD},
    "cupons": {"nome": "Promocao", "codigo": "PROMO", "data_validade": "2099-12-31",
               "tipo_desconto": "PERCENTUAL", "valor_desconto": "10"},
    "formas_pagamento": {"nome": "Pagamento importado"},
    "categorias_financeiras": {"nome": "Categoria importada", "tipo_categoria": "SAIDA"},
}


@pytest.mark.parametrize("entity", LAYOUT_ROWS)
def test_all_six_layouts_preview_import_export_reimport(transaction_app, entity):
    with transaction_app.app_context():
        rows = [LAYOUT_ROWS[entity]]
        preview = import_rows(entity, rows, pre_validar=True)
        assert preview["validas"] == 1 and not preview["confirmado"] and preview["criadas"] == 0
        result = import_rows(entity, rows)
        assert result["confirmado"] and result["criadas"] == 1, result
        template = ImportExportService.gerar_template(entity, 1, scope())
        assert "instrucoes" in load_workbook(BytesIO(template["content"])).sheetnames
        exported = ImportExportService.exportar_entidade(entity, 1, scope())
        workbook = load_workbook(BytesIO(exported["content"]), data_only=True)
        sheet_rows = list(workbook["dados"].iter_rows(values_only=True))
        assert len(sheet_rows) >= 2
        replay = ImportExportService.importar_entidade(entity,
            FileStorage(stream=BytesIO(exported["content"]), filename="reimport.xlsx"), 1, scope(), 1)
        assert replay["confirmado"] and replay["criadas"] == 0, replay
        if entity == "produtos":
            record = ProdutoEmpresa.query.one()
            assert record.valor_varejo == 10 and record.valor_atacado == 8
            assert record.quantidade_minima_atacado == 5 and record.estoque_atual == 3
            assert MovimentoEstoque.query.count() == 1
        if entity == "funcionarios":
            headers = sheet_rows[0]
            assert all(not row[headers.index("senha")] for row in sheet_rows[1:])


@pytest.mark.parametrize("entity", LAYOUT_ROWS)
def test_all_layouts_reject_entire_batch_when_last_row_invalid(transaction_app, entity):
    with transaction_app.app_context():
        row = LAYOUT_ROWS[entity]
        result = import_rows(entity, [row, {**row, "nome": "", "ativo": "SIM"}])
        assert result["falhas"] == 1 and result["sucesso"] == 0 and not result["confirmado"], result
        assert result["erros"][0]["linha"] == 3
        retry = import_rows(entity, [row])
        assert retry["criadas"] == 1, retry


def test_import_stock_update_and_conflicting_barcode_rollback(transaction_app):
    with transaction_app.app_context():
        first = product("First", codigo_barras="123")
        product("Second", codigo_barras="456")
        assert not import_rows("produtos", [{"empresa": "Empresa 1-1", "nome": "First", "codigo_barras": "456"}])["confirmado"]
        assert not import_rows("produtos", [{"empresa": "Empresa 1-1", "nome": "First", "estoque_atual": 4}])["confirmado"]
        assert first.produto.codigo_barras == "123" and first.estoque_atual == 0


def test_concurrent_import_last_quota_and_no_orphans(transaction_app):
    with transaction_app.app_context():
        db.session.get(Tenant, 1).limite_produtos = 1
        db.session.commit()
    results = parallel(transaction_app, lambda index: import_rows("produtos", [{
        "empresa": "Empresa 1-1", "nome": f"Import {index}", "categoria": f"Category {index}"}]))
    assert sum(result["confirmado"] for result in results) == 1
    with transaction_app.app_context():
        assert Produto.query.count() == ProdutoEmpresa.query.count() == CategoriaProduto.query.count() == 1


def test_http_preview_and_multicompany_employee(transaction_app):
    client = transaction_app.test_client()
    assert login(client).status_code == 302
    with transaction_app.app_context():
        payload = employee()
    result = client.post("/api/funcionarios/", json=payload, headers=csrf_headers(client))
    assert result.status_code == 201, result.json
    exported = upload("categorias", [{"nome": "Preview"}])
    response = client.post("/api/importacao-exportacao/importar", data={"entidade": "categorias",
        "pre_validar": "true", "arquivo": (exported.stream, "preview.xlsx")}, headers=csrf_headers(client))
    assert response.status_code == 200 and not response.json["data"]["confirmado"]
    with transaction_app.app_context():
        assert CategoriaProduto.query.count() == 0


def test_role_replacement_failure_preserves_permissions_and_name(transaction_app):
    with transaction_app.app_context():
        permission_id = RolePermission.query.filter_by(tenant_id=1).first().permission_id
        role = RoleService.criar({"nome": "Perfil", "codigo": "perfil", "permission_ids": [permission_id]}, 1)
        role_id = role.id

        def fail(*args):
            raise RuntimeError("replacement failure")

        event.listen(RolePermission, "after_insert", fail)
        try:
            with pytest.raises(RuntimeError):
                RoleService.atualizar(role_id, {"nome": "Changed", "codigo": "changed", "permission_ids": [permission_id]}, 1)
        finally:
            event.remove(RolePermission, "after_insert", fail)
        assert db.session.get(Role, role_id).nome == "Perfil"
        assert RolePermission.query.filter_by(role_id=role_id).one().permission_id == permission_id


def test_cancel_failure_restores_sale_stock_finance_and_idempotency(transaction_app, monkeypatch):
    from app.services.financeiro_service import FinanceiroService
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        sale = PdvService.criar_venda(sale_payload(product_id), 1, scope(), 1)
        financial_count = LancamentoFinanceiro.query.count()

        def fail(**kwargs):
            db.session.flush()
            raise RuntimeError("cancel failure")

        monkeypatch.setattr(FinanceiroService, "registrar_estorno_da_venda", fail)
        with pytest.raises(RuntimeError):
            PdvService.cancelar_venda(sale["id"], {"idempotency_key": "cancel"}, 1, scope(), 1)
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 9
        assert db.session.get(Venda, sale["id"]).status.value == "FINALIZADA"
        assert MovimentoEstoque.query.count() == OperacaoIdempotente.query.count() == 1
        assert LancamentoFinanceiro.query.count() == financial_count


def test_commit_failure_does_not_send_notification_or_leave_sale(transaction_app, monkeypatch):
    from app.services.cliente_service import ClienteService
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        payload = sale_payload(product_id)
        calls = []
        monkeypatch.setattr(ClienteService, "enviar_email_venda_automatica", lambda **kwargs: calls.append(kwargs))

        def fail():
            raise RuntimeError("commit failure")

        with monkeypatch.context() as commit_patch:
            commit_patch.setattr(db.session, "commit", fail)
            with pytest.raises(RuntimeError):
                PdvService.criar_venda(payload, 1, scope(), 1)
        assert calls == []
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 10
        assert Venda.query.count() == MovimentoEstoque.query.count() == OperacaoIdempotente.query.count() == 0


def test_uncancelled_sale_cannot_restore_stock(transaction_app):
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        sale = PdvService.criar_venda(sale_payload(product_id), 1, scope(), 1)
        with pytest.raises(ValueError, match="nao cancelado"):
            EstoqueService.registrar_entrada_por_cancelamento_venda(sale["id"], 1,
                [{"item_venda_id": sale["itens"][0]["id"], "produto_id": product_id, "quantidade": 1}], 1, 1, scope())
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 9


def test_manual_cancellation_concurrent_replay(transaction_app):
    with transaction_app.app_context():
        record_id, _ = stock_setup()
        movement = EstoqueService.registrar_movimentacao_manual({"produto_empresa_id": record_id,
            "tipo_movimento": "SAIDA", "motivo": "AJUSTE", "quantidade": 1}, 1, scope(), 1)
        movement_id = movement.id
    results = parallel(transaction_app, lambda index: EstoqueService.cancelar_movimento(movement_id, {}, 1, scope(), 1).id, 3)
    assert len(set(results)) == 1
    with transaction_app.app_context():
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 10
        assert MovimentoEstoque.query.count() == 2


def test_deactivation_is_local_to_product_company(transaction_app):
    with transaction_app.app_context():
        first = product()
        first_id, product_id = first.id, first.produto_id
        second = ProdutoService.criar({"produto_id": product_id, "empresa_id": 2}, 1, scope(), 1)
        second_id = second.id
        ProdutoService.atualizar(first_id, {"nome": "Produto", "empresa_id": 1, "ativo": False}, 1, scope())
        assert not db.session.get(ProdutoEmpresa, first_id).ativo
        assert db.session.get(ProdutoEmpresa, second_id).ativo
        assert db.session.get(Produto, product_id).ativo


@pytest.mark.parametrize("row", [
    {"nome": "=1+1"}, {"nome": "Invalid", "valor_atacado": "11", "valor_varejo": "10"},
    {"nome": "Invalid", "quantidade_minima_atacado": "1.5"}, {"nome": "Invalid", "ncm": "123"},
])
def test_import_invalid_values_rollback(transaction_app, row):
    with transaction_app.app_context():
        result = import_rows("produtos", [{"empresa": "Empresa 1-1", **row}])
        assert result["falhas"] == 1 and not result["confirmado"]
        assert Produto.query.count() == ProdutoEmpresa.query.count() == 0


def test_export_formula_like_text_is_literal(transaction_app):
    with transaction_app.app_context():
        CategoriaService.criar({"nome": "=1+1"}, 1)
        exported = ImportExportService.exportar_entidade("categorias", 1, scope())
        workbook = load_workbook(BytesIO(exported["content"]), data_only=False)
        assert workbook["dados"]["A2"].data_type == "s"
        replay = ImportExportService.importar_entidade("categorias",
            FileStorage(stream=BytesIO(exported["content"]), filename="literal.xlsx"), 1, scope(), 1)
        assert replay["confirmado"] and replay["atualizadas"] == 1


def test_http_sale_requires_key_and_replays_same_business_response(transaction_app):
    client = transaction_app.test_client()
    assert login(client).status_code == 302
    with transaction_app.app_context():
        record_id, product_id = stock_setup()
        payload = sale_payload(product_id, quantity=2)
    without_key = {key: value for key, value in payload.items() if key != "idempotency_key"}
    assert client.post("/api/pdv/vendas", json=without_key, headers=csrf_headers(client)).status_code == 400
    first = client.post("/api/pdv/vendas", json=payload, headers=csrf_headers(client))
    second = client.post("/api/pdv/vendas", json=payload, headers=csrf_headers(client))
    assert first.status_code == second.status_code == 201, (first.json, second.json)
    assert first.json["data"] == second.json["data"]
    sale = first.json["data"]
    endpoint = f"/api/pdv/vendas/{sale['id']}/itens/{sale['itens'][0]['id']}/cancelar"
    assert client.post(endpoint, json={}, headers=csrf_headers(client)).status_code == 400
    partial = {"idempotency_key": "http-partial", "quantidade": 1}
    assert client.post(endpoint, json=partial, headers=csrf_headers(client)).status_code == 200
    assert client.post(endpoint, json=partial, headers=csrf_headers(client)).status_code == 200
    with transaction_app.app_context():
        assert db.session.get(ProdutoEmpresa, record_id).estoque_atual == 9
        assert Venda.query.count() == 1 and MovimentoEstoque.query.count() == 2


def test_migration_rejects_duplicate_legacy_sale_debits_atomically(transaction_app):
    from flask_migrate import downgrade, upgrade
    from sqlalchemy.exc import IntegrityError
    with transaction_app.app_context():
        _, product_id = stock_setup()
        PdvService.criar_venda(sale_payload(product_id), 1, scope(), 1)
        db.session.remove()
        downgrade(revision="4d5e6f7a8b9c")
        with db.engine.begin() as connection:
            duplicate_id = connection.execute(text("""
                INSERT INTO movimentos_estoque(tenant_id,empresa_id,produto_id,venda_id,item_venda_id,
                    tipo_movimento,motivo,quantidade,revertido,data_movimento,criado_em,atualizado_em)
                SELECT tenant_id,empresa_id,produto_id,venda_id,item_venda_id,tipo_movimento,motivo,
                    quantidade,revertido,data_movimento,criado_em,atualizado_em FROM movimentos_estoque LIMIT 1
                RETURNING id
            """)).scalar_one()
        try:
            with pytest.raises(IntegrityError):
                upgrade()
            with db.engine.connect() as connection:
                assert connection.execute(text("SELECT version_num FROM alembic_version")).scalar_one() == "4d5e6f7a8b9c"
                assert connection.execute(text("SELECT to_regclass('operacoes_idempotentes')")).scalar_one() is None
        finally:
            with db.engine.begin() as connection:
                connection.execute(text("DELETE FROM movimentos_estoque WHERE id=:duplicate"), {"duplicate": duplicate_id})
            upgrade()
